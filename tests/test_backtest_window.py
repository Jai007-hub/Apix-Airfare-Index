from datetime import date

from db.models import IndexValue
from validation.backtest import (
    deviation_profile,
    directional_agreement,
    error_metrics,
    run_backtest,
)


def _seed_monthly_index(db_session, months, base=date(2025, 1, 1)):
    for (year, month), value in months.items():
        db_session.add(
            IndexValue(
                frequency="monthly",
                period_date=date(year, month, 1),
                apix_value=value,
                base_period_date=base,
            )
        )
    db_session.commit()


def test_days_covered_spans_whole_months_not_just_month_starts(db_session, tmp_path):
    """The requirement is written in days but CPI only publishes monthly, so the
    window has to be measured to the END of the last compared month."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CPI Data"
    ws.append(
        ["base_year", "series", "year", "month", "state", "sector", "division", "group",
         "class", "sub_class", "item", "code", "index", "inflation", "imputation"]
    )
    # Three consecutive months of All-India CPI: Jan, Feb, Mar 2025.
    for month_name_, idx in [("January", 100.0), ("February", 102.0), ("March", 104.0)]:
        ws.append(
            ["2024", "Current", "2025", month_name_, "All India", "Combined", "Transport",
             "grp", "cls", "sub", "Airfare", "07", idx, 0.0, "N"]
        )
    xlsx = tmp_path / "cpi.xlsx"
    wb.save(xlsx)

    _seed_monthly_index(
        db_session, {(2025, 1): 100.0, (2025, 2): 101.0, (2025, 3): 103.0}
    )

    result = run_backtest(db_session, str(xlsx))

    assert result["n_months_compared"] == 3
    assert result["window_start"] == date(2025, 1, 1)
    # Must run to 31 March, not 1 March -- otherwise the day count is short by
    # nearly a month and understates the coverage.
    assert result["window_end"] == date(2025, 3, 31)
    assert result["days_covered"] == 31 + 28 + 31  # Jan + Feb + Mar 2025


def test_days_covered_handles_a_december_end_month(db_session, tmp_path):
    """December has to roll into the next year when finding the month end."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CPI Data"
    ws.append(["c"] * 15)
    for month_name_, idx in [("November", 100.0), ("December", 101.0)]:
        ws.append(
            ["2024", "Current", "2025", month_name_, "All India", "Combined", "Transport",
             "grp", "cls", "sub", "Airfare", "07", idx, 0.0, "N"]
        )
    xlsx = tmp_path / "cpi.xlsx"
    wb.save(xlsx)

    _seed_monthly_index(db_session, {(2025, 11): 100.0, (2025, 12): 102.0})

    result = run_backtest(db_session, str(xlsx))

    assert result["window_end"] == date(2025, 12, 31)
    assert result["days_covered"] == 30 + 31  # Nov + Dec


def test_window_comfortably_exceeds_the_30_day_requirement(db_session, tmp_path):
    """Two adjacent months already clear the problem statement's 30-day floor."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CPI Data"
    ws.append(["c"] * 15)
    for month_name_, idx in [("June", 120.0), ("July", 121.0)]:
        ws.append(
            ["2024", "Current", "2026", month_name_, "All India", "Combined", "Transport",
             "grp", "cls", "sub", "Airfare", "07", idx, 0.0, "N"]
        )
    xlsx = tmp_path / "cpi.xlsx"
    wb.save(xlsx)

    _seed_monthly_index(db_session, {(2026, 6): 118.0, (2026, 7): 119.0})

    assert run_backtest(db_session, str(xlsx))["days_covered"] >= 30


# -- directional agreement ----------------------------------------------------
# Pearson r on a short series swings on one noisy month; "did both move the
# same way" is the blunter check shown alongside it.


def test_series_that_always_move_together_agree_completely():
    apix = [100.0, 105.0, 103.0, 108.0]
    cpi = [200.0, 210.0, 190.0, 220.0]  # same ups and downs, different scale
    assert directional_agreement(apix, cpi) == {"matches": 3, "comparisons": 3, "pct": 100.0}


def test_series_that_always_move_oppositely_agree_never():
    apix = [100.0, 105.0, 103.0, 108.0]
    cpi = [200.0, 190.0, 210.0, 190.0]
    assert directional_agreement(apix, cpi) == {"matches": 0, "comparisons": 3, "pct": 0.0}


def test_denominator_is_moves_not_months():
    """Four months give three month-on-month moves -- reporting "out of 4"
    would overstate the sample."""
    result = directional_agreement([1.0, 2.0, 3.0, 4.0], [1.0, 2.0, 3.0, 4.0])
    assert result["comparisons"] == 3


def test_mixed_directions_are_counted_not_rounded_away():
    #        up/up      down/up     up/up
    apix = [100.0, 110.0, 105.0, 115.0]
    cpi = [100.0, 110.0, 120.0, 130.0]
    assert directional_agreement(apix, cpi) == {"matches": 2, "comparisons": 3, "pct": 66.7}


def test_a_flat_month_only_agrees_with_another_flat_month():
    assert directional_agreement([100.0, 100.0], [100.0, 100.0])["matches"] == 1
    assert directional_agreement([100.0, 100.0], [100.0, 101.0])["matches"] == 0


def test_too_few_points_reports_nothing_rather_than_a_misleading_zero():
    assert directional_agreement([100.0], [100.0]) is None
    assert directional_agreement([], []) is None
    # Mismatched lengths would silently truncate; refuse instead.
    assert directional_agreement([100.0, 101.0], [100.0]) is None


# -- deviation profile --------------------------------------------------------
# The distribution behind the MAPE headline, so a mean isn't read on its own.


def _pt(year, month, pct):
    return {"year": year, "month": month, "pct_diff": pct}


def test_rebase_anchor_is_not_reported_as_the_best_month():
    """Rebasing forces the first overlapping month to 0.000% deviation. That
    is an arithmetic identity, not accuracy, so it must not win 'best month'
    -- otherwise every back-test claims a perfect month it did not earn."""
    points = [_pt(2025, 1, 0.0), _pt(2025, 2, -3.0), _pt(2025, 3, 8.0)]

    profile = deviation_profile(points)

    assert profile["best_month"] == {"year": 2025, "month": 2, "abs_pct": 3.0}
    assert profile["n"] == 2
    assert profile["excludes_rebase_anchor"] is True


def test_reports_the_worst_month_by_absolute_deviation():
    """Sign must not decide it -- a -19% month is as far off as +19%."""
    points = [_pt(2025, 1, 0.0), _pt(2025, 2, 4.0), _pt(2025, 3, -19.0), _pt(2025, 4, 12.0)]

    assert deviation_profile(points)["worst_month"] == {
        "year": 2025,
        "month": 3,
        "abs_pct": 19.0,
    }


def test_median_describes_the_middle_not_the_mean():
    """One terrible month drags the mean but not the median -- which is the
    whole reason both are shown."""
    points = [_pt(2025, 1, 0.0), _pt(2025, 2, 2.0), _pt(2025, 3, 3.0), _pt(2025, 4, 40.0)]

    profile = deviation_profile(points)

    assert profile["median_abs_pct"] == 3.0  # mean of the same three is 15.0


def test_within_5pct_includes_the_boundary_and_ignores_sign():
    points = [
        _pt(2025, 1, 0.0),
        _pt(2025, 2, 5.0),
        _pt(2025, 3, -5.0),
        _pt(2025, 4, 5.01),
    ]

    assert deviation_profile(points)["within_5pct"] == 2


def test_nothing_to_profile_once_the_anchor_is_removed():
    assert deviation_profile([_pt(2025, 1, 0.0)]) is None
    assert deviation_profile([]) is None


# -- error metrics ------------------------------------------------------------
# The regression measures a judge expects to see: MAE, MSE, RMSE and bias.


def test_error_metrics_on_a_hand_computable_series():
    apix = [102.0, 98.0, 105.0]
    cpi = [100.0, 100.0, 100.0]
    # errors +2, -2, +5 -> MAE 3, MSE 11, RMSE sqrt(11), bias +5/3
    m = error_metrics(apix, cpi)

    assert m["mae_index_points"] == 3.0
    assert m["mse_index_points"] == 11.0
    assert m["rmse_index_points"] == round(11 ** 0.5, 3)
    assert m["mean_bias_index_points"] == 1.667


def test_rmse_exceeds_mae_when_the_error_is_concentrated():
    """RMSE >> MAE is the signal that a few bad months dominate, which is
    exactly what it is reported for."""
    spread = error_metrics([101.0, 101.0, 101.0, 101.0], [100.0] * 4)
    spiky = error_metrics([100.0, 100.0, 100.0, 104.0], [100.0] * 4)

    assert spread["mae_index_points"] == spiky["mae_index_points"] == 1.0
    assert spread["rmse_index_points"] == 1.0
    assert spiky["rmse_index_points"] > spread["rmse_index_points"]


def test_bias_separates_wandering_from_consistently_off():
    """MAE says how far off; bias says which side. Both series below are 2
    points out on average, but one sits high throughout and the other
    straddles."""
    high = error_metrics([102.0, 102.0], [100.0, 100.0])
    straddles = error_metrics([102.0, 98.0], [100.0, 100.0])

    assert high["mae_index_points"] == straddles["mae_index_points"] == 2.0
    assert high["mean_bias_index_points"] == 2.0
    assert straddles["mean_bias_index_points"] == 0.0


def test_mismatched_or_empty_series_report_nothing():
    assert error_metrics([], []) is None
    assert error_metrics([100.0, 101.0], [100.0]) is None
